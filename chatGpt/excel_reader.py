from openpyxl import load_workbook


def read_test_cases(path):
    workbook = load_workbook(path)
    sheet = workbook.active

    test_cases = {}
    for case_id, case_name, url, step_no, step_text in sheet.iter_rows(min_row=2, values_only=True):
        if case_id not in test_cases:
            test_cases[case_id] = {
                "id": case_id,
                "name": case_name,
                "url": url,
                "steps": []
            }
        test_cases[case_id]["steps"].append({"number": step_no, "text": step_text})
    workbook.close()
    return list(test_cases.values())


       
